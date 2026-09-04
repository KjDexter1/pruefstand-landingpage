# -*- coding: utf-8 -*-
"""Traegt den Ausnahmenblock auf Blatt 4 nach.

Blatt 2 wird nicht mehr hier veraendert -- die Punkte stehen jetzt
direkt in baue_arbeitsmappe.py. Nachtraegliches insert_rows verschiebt
verbundene Zellen nicht mit und zerstoert dabei Inhalte.
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

GRUEN="0E5B50"; ROST="A93C1B"; GELB="FFF6D6"; LINIE="D4D2C9"
def f(size=9, bold=False, color="1A1D1B"):
    return Font(name="Arial", size=size, bold=bold, color=color)
kante = Side(style="thin", color=LINIE)
rahmen = Border(left=kante, right=kante, top=kante, bottom=kante)
oben = Alignment(vertical="top", horizontal="left", wrap_text=True)

wb = load_workbook("E-Rechnung_Arbeitsmappe.xlsx")

# ================= Blatt 4: eigener Ausnahmenblock ==================
ws4 = wb["4 Angebot & Regeln"]
r = ws4.max_row + 2

ws4.cell(row=r, column=1,
         value="Ausnahmen — wer die Rechnung NICHT elektronisch stellen muss").font = \
    f(12, bold=True, color=GRUEN)
r += 1
ws4.cell(row=r, column=1,
         value="Vor jedem Angebot pruefen. Wer hier hineinfaellt, hat ein kleineres "
               "Problem als geschildert — das offen zu sagen kostet einen Auftrag "
               "und bringt Glaubwuerdigkeit.").font = f(9, color="55605A")
ws4.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws4.row_dimensions[r].height = 26
r += 2

kopf = ["Ausnahme", "Was das bedeutet", "Folge fuers Angebot", ""]
for i, text in enumerate(kopf, start=1):
    c = ws4.cell(row=r, column=i, value=text)
    c.font = f(9, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=GRUEN)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    c.border = rahmen
ws4.row_dimensions[r].height = 30
r += 1

ausnahmen = [
    ("Kleinbetragsrechnungen bis 250 € brutto",
     "Rechnungen bis 250 € brutto sind von der Ausstellungspflicht ausgenommen. "
     "Der Betrieb darf sie weiterhin als Papier oder PDF stellen.",
     "Wer fast nur Kleinbetraege abrechnet, braucht kaum etwas. Ehrlich sagen. "
     "Wer beides hat, braucht die Umstellung trotzdem — dann fuer den Rest."),
    ("Rechnungen an Privatpersonen (B2C)",
     "Die Pflicht gilt nur zwischen inlaendischen Unternehmen. An Privatkunden "
     "gibt es keine E-Rechnungspflicht.",
     "Ein reiner Privatkundenbetrieb ist nicht betroffen. Mischbetriebe schon — "
     "aber nur fuer den gewerblichen Teil."),
    ("Bestimmte steuerfreie Umsaetze nach § 4 UStG",
     "Einzelne nach § 4 UStG steuerfreie Leistungen sind ausgenommen. "
     "Welche genau, entscheidet der Steuerberater.",
     "Nicht selbst beurteilen. An den Steuerberater verweisen — das ist "
     "Steuerrecht, nicht Technik."),
    ("Kleinunternehmer nach § 19 UStG",
     "Dauerhaft von der Ausstellungspflicht befreit. Empfangen muessen sie "
     "trotzdem, und zwar schon seit 2025.",
     "Kein Versandpaket anbieten. Stattdessen das kleine Empfangspaket — "
     "die Pflicht laeuft bereits."),
]
for name, bedeutung, folge in ausnahmen:
    ws4.cell(row=r, column=1, value=name).font = f(10, bold=True)
    ws4.cell(row=r, column=2, value=bedeutung).font = f(9)
    ws4.cell(row=r, column=3, value=folge).font = f(9)
    for s in range(1, 5):
        ws4.cell(row=r, column=s).alignment = oben
        ws4.cell(row=r, column=s).border = rahmen
    ws4.row_dimensions[r].height = 44
    r += 1

r += 1
c = ws4.cell(row=r, column=1,
             value="Die Frage am Telefon dafuer: „Wie hoch ist bei Ihnen eine "
                   "durchschnittliche Rechnung, und geht die an Betriebe oder an "
                   "Privatleute?“ — beantwortet beide Ausnahmen auf einmal.")
c.font = f(10, bold=True, color=ROST)
c.alignment = oben
ws4.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws4.row_dimensions[r].height = 32

wb.save("E-Rechnung_Arbeitsmappe.xlsx")
print("nachgetragen")
