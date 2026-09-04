# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

WEISS="FFFFFF"; GELB="FFF6D6"; ROST="A93C1B"
def f(size=9, bold=False, color="1A1D1B"):
    return Font(name="Arial", size=size, bold=bold, color=color)
oben = Alignment(vertical="top", wrap_text=True)
OFFEN = "— noch offen —"
P13B = OFFEN + " — WICHTIGSTE FRAGE, siehe Blatt 3"

wb = load_workbook("E-Rechnung_Arbeitsmappe.xlsx")
ws = wb["1 Systeme"]

daten = {
 9: {  # WISO Mein Buero
  "D":"04.09.2026","E":OFFEN,"F":"ja","G":"ZUGFeRD automatisch, XRechnung als XML-Download",
  "H":"Einstellungen > Meine Firma > E-Rechnungen",
  "I":"Bankverbindung und Kontaktdaten pflicht. Nur B2B — an Privatkunden keine E-Rechnung.",
  "J":P13B,"K":OFFEN,"L":OFFEN,
  "M":"Jede Rechnung wird automatisch als ZUGFeRD erzeugt; die XRechnung-XML muss man nach dem Abschliessen separat herunterladen. Desktop- und Web-Version unterscheiden sich — vorher klaeren, welche der Kunde hat.",
  "N":None,"O":"handbuch.meinbuero.de/mbr-e-rechnung/e-rechnung-erstellen",
 },
 11: { # FastBill
  "D":"04.09.2026","E":OFFEN,"F":"ja","G":"XRechnung + ZUGFeRD",
  "H":"E-Rechnungs-Option aktivieren, dann wird bei „Rechnung erstellen“ automatisch beides erzeugt",
  "I":"Alle Pflichtangaben muessen gefuellt sein, sonst wird keine E-Rechnung erzeugt — welche genau: offen",
  "J":P13B,"K":OFFEN,"L":OFFEN,
  "M":"Erzeugt ZUGFeRD und XRechnung gleichzeitig. Erzeugt sie nur, wenn alle Pflichtfelder gefuellt sind — sonst still kein E-Rechnungs-Ergebnis. Genau hinsehen.",
  "N":None,"O":"support.fastbill.com/hc/de/articles/22209126147484",
 },
 12: { # Billomat
  "D":"04.09.2026","E":OFFEN,"F":"ja","G":"XRechnung + ZUGFeRD",
  "H":OFFEN,"I":OFFEN,"J":P13B,"K":OFFEN,"L":OFFEN,
  "M":"Unterstuetzung bestaetigt, Einrichtungsdetails noch nicht recherchiert.",
  "N":None,"O":"—",
 },
 13: { # Papierkram
  "D":"04.09.2026","E":OFFEN,"F":OFFEN,"G":OFFEN,"H":OFFEN,"I":OFFEN,"J":P13B,"K":OFFEN,"L":OFFEN,
  "M":"Nichts Belastbares gefunden. Herstellerhilfe direkt aufrufen.",
  "N":None,"O":"—",
 },
 14: { # SumUp
  "D":"04.09.2026","E":OFFEN,"F":OFFEN,"G":OFFEN,"H":OFFEN,"I":OFFEN,"J":P13B,"K":OFFEN,"L":OFFEN,
  "M":"Nichts Belastbares gefunden. Eher Kleinstgewerbe — pruefen, ob die Zielgruppe ueberhaupt zahlt.",
  "N":None,"O":"—",
 },
 15: { # Collmex
  "D":"04.09.2026","E":OFFEN,"F":OFFEN,"G":OFFEN,"H":OFFEN,"I":OFFEN,"J":P13B,"K":OFFEN,"L":OFFEN,
  "M":"Nichts Belastbares gefunden. Collmex ist im Handwerk verbreitet — lohnt eigene Recherche.",
  "N":None,"O":"—",
 },
}

for zeile, felder in daten.items():
    for spalte, wert in felder.items():
        if wert is None: continue
        c = ws[f"{spalte}{zeile}"]
        c.value = wert
        offen = isinstance(wert,str) and OFFEN in wert
        c.fill = PatternFill("solid", fgColor=GELB if offen else WEISS)
        c.font = f(9, bold=(spalte=="J"), color=(ROST if spalte=="J" else "1A1D1B"))
        c.alignment = oben
    ws.row_dimensions[zeile].height = 74

# --- Fehlerkatalog: exakte Codes fuer § 13b nachtragen -------------
ws3 = wb["3 Fehlerkatalog"]
ws3["D6"] = ("Vier Angaben muessen zusammenpassen: Steuerkategorie „AE“, Steuersatz 0, "
             "Befreiungsgrund-Code „VATEX-EU-AE“ und der Befreiungstext "
             "„Steuerschuldnerschaft des Leistungsempfaengers“. Pruefen, ob die Software das "
             "automatisch setzt, sobald der Kunde als §13b-pflichtig markiert ist. "
             "ACHTUNG: Lexware Office kann §13b derzeit gar nicht als E-Rechnung ausgeben — "
             "siehe Blatt 1. Ob §13b gilt, entscheidet der Steuerberater, nicht du.")
ws3["D6"].font = f(9, bold=True, color=ROST)
ws3["D6"].alignment = oben
ws3.row_dimensions[6].height = 90

# --- Blatt 4: ehrlicher Hinweis zum Wettbewerb ---------------------
ws4 = wb["4 Angebot & Regeln"]
r = ws4.max_row + 2
ws4.cell(row=r, column=1, value="Zum Leser auf der Website — ehrlich eingeordnet").font = f(11, bold=True, color=ROST)
r += 1
for satz in [
  "Kostenlose E-Rechnungs-Betrachter gibt es bereits: Quba Viewer (vom Wirtschaftsministerium empfohlen), "
  "RIB Viewer, Mustang. Dein Leser ist technisch nichts Einzigartiges.",
  "Sein Wert liegt nicht in der Funktion, sondern darin, dass er auf DEINER Seite steht und dort "
  "das Gespraech eroeffnet. Verkaufe nie den Betrachter — verkaufe die Einrichtung.",
]:
    c = ws4.cell(row=r, column=1, value="•  " + satz)
    c.font = f(9); c.alignment = Alignment(vertical="top", wrap_text=True)
    ws4.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws4.row_dimensions[r].height = 32
    r += 1

wb.save("E-Rechnung_Arbeitsmappe.xlsx")
print("fertig")
